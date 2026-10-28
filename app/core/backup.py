import hashlib
import io
import json
import sqlite3
import zipfile
from app.core.db import Database
from datetime import datetime, timezone
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook = None
    load_workbook = None

SHEETS = {
    "cashier": """
        SELECT transaction_id,created_at,hardware_id,member_id,credential_id,account_id,
               'DEPOSIT' AS transaction_type,amount,balance_before,balance_after,status
        FROM cashier_deposits
        ORDER BY created_at
    """,
    "vending": """
        SELECT transaction_id,created_at,machine_id,hardware_id,credential_id,account_id,
               'PURCHASE' AS transaction_type,product_id,amount,payment_status,
               dispense_status,status
        FROM vending_transactions v
        LEFT JOIN vending_machines m ON m.machine_id=v.machine_id
        ORDER BY v.created_at
    """,
    "payments": """
        SELECT transaction_id,created_at,transaction_source,credential_id,
               account_id,method,provider,original_amount,discount_amount,final_amount,
               amount,status,idempotency_key
        FROM payment_transactions
        ORDER BY created_at
    """,
    "ledger": """
        SELECT id,created_at,account_id,direction,amount,reference
        FROM financial_ledger
        ORDER BY created_at
    """,
    "audit": """
        SELECT id,created_at,event_type,entity_type,entity_id,detail
        FROM audit_events
        ORDER BY created_at
    """,
}

def _rows(db, sql):
    with db.connect() as c:
        cur=c.execute(sql)
        cols=[d[0] for d in cur.description]
        return cols,[tuple(r) for r in cur.fetchall()]

def export_excel(db, target, start=None, end=None):
    if Workbook is None:
        raise RuntimeError("openpyxl_required")
    target=Path(target)
    target.parent.mkdir(parents=True,exist_ok=True)
    wb=Workbook()
    wb.remove(wb.active)
    for name,sql in SHEETS.items():
        args=()
        if start:
            # Filter the selected result columns in an outer query. This avoids
            # ambiguous created_at references in joined exports (notably vending,
            # which joins vending_transactions and vending_machines).
            order_by = sql.upper().rfind("ORDER BY")
            base_sql = sql[:order_by].rstrip() if order_by >= 0 else sql.rstrip().rstrip(";")
            sql = (
                "SELECT * FROM (" + base_sql + ") AS export_rows "
                "WHERE created_at>=? AND created_at<? ORDER BY created_at"
            )
            args=(start,end)
        with db.connect() as c:
            cur=c.execute(sql,args)
            headers=[d[0] for d in cur.description]
            rows=cur.fetchall()
        ws=wb.create_sheet(name[:31])
        ws.append(headers)
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes="A2"
        ws.auto_filter.ref=ws.dimensions
    wb.save(target)
    return target

def create_native_backup(db, target, include_excel=True):
    target=Path(target)
    target.parent.mkdir(parents=True,exist_ok=True)
    with db.connect() as c:
        # Keep the schema/data dump on one SQLite read snapshot. Otherwise a
        # concurrent transaction could commit between table dumps and produce
        # an archive whose tables reflect different points in time.
        c.execute("BEGIN")
        integrity = c.execute("PRAGMA integrity_check").fetchone()[0]
        if integrity != "ok":
            raise ValueError("backup_database_integrity_check_failed")
        foreign_key_errors = c.execute("PRAGMA foreign_key_check").fetchall()
        if foreign_key_errors:
            raise ValueError("backup_database_foreign_key_check_failed")
        raw=c.execute("SELECT name,sql FROM sqlite_master WHERE type='table' ORDER BY name").fetchall()
        dump="\n".join(c.iterdump()).encode()
    manifest={
        "format":"M-Group-native-backup-v1",
        "created_at":datetime.now(timezone.utc).isoformat(),
        "retention_months":12,
        "tables":[r[0] for r in raw],
    }
    with zipfile.ZipFile(target,"w",zipfile.ZIP_DEFLATED) as z:
        z.writestr("database.sql",dump)
        z.writestr("manifest.json",json.dumps(manifest,indent=2))
        if include_excel:
            if Workbook is None:
                raise RuntimeError("openpyxl_required")
            buf=target.with_suffix(".xlsx")
            export_excel(db,buf)
            z.write(buf,"excel/M-Group-export.xlsx")
            buf.unlink(missing_ok=True)
        digest=hashlib.sha256(dump).hexdigest()
        z.writestr("checksum.sha256",digest+"  database.sql\n")
    return target

def restore_native_backup(db, archive):
    archive = Path(archive)
    max_archive_bytes = 100 * 1024 * 1024
    max_expanded_bytes = 250 * 1024 * 1024
    if archive.stat().st_size > max_archive_bytes:
        raise ValueError("backup_archive_too_large")

    with zipfile.ZipFile(archive) as z:
        infos = z.infolist()
        names = [info.filename for info in infos]
        required = {"manifest.json", "database.sql", "checksum.sha256"}
        if len(names) != len(set(names)) or not required.issubset(names):
            raise ValueError("invalid_backup_archive")
        if len(infos) > 1000 or sum(info.file_size for info in infos) > max_expanded_bytes:
            raise ValueError("backup_archive_too_large")
        info_by_name = {info.filename: info for info in infos}
        if info_by_name["manifest.json"].file_size > 1024 * 1024:
            raise ValueError("backup_manifest_too_large")
        if info_by_name["database.sql"].file_size > 200 * 1024 * 1024:
            raise ValueError("backup_database_too_large")
        manifest = json.loads(z.read("manifest.json"))
        if not isinstance(manifest, dict) or manifest.get("format") != "M-Group-native-backup-v1":
            raise ValueError("unsupported_backup_format")
        dump = z.read("database.sql")
        checksum_parts = z.read("checksum.sha256").decode("ascii").split()
        if not checksum_parts or len(checksum_parts[0]) != 64:
            raise ValueError("invalid_backup_checksum")
        expected = checksum_parts[0]
        if hashlib.sha256(dump).hexdigest() != expected:
            raise ValueError("backup_checksum_mismatch")
        sql = dump.decode("utf-8")

    # Validate, migrate and integrity-check a temporary copy first. Never replace
    # the active database until the complete restore path is known to be usable.
    tmp = archive.with_suffix(".restore.db")
    tmp.unlink(missing_ok=True)
    probe = None
    try:
        probe = sqlite3.connect(tmp)
        # sqlite3.iterdump may emit child-table INSERTs before parent-table
        # definitions/data. Disable FK enforcement only during archive loading,
        # then validate every relationship after migrations below.
        probe.execute("PRAGMA foreign_keys=OFF")
        probe.executescript(sql)
        probe.commit()
        probe.execute("PRAGMA foreign_keys=ON")
        probe.close()
        probe = None

        restored_db = Database(tmp)
        restored_db.migrate()
        with restored_db.connect() as check:
            admin_count=check.execute("SELECT COUNT(*) FROM admin_users").fetchone()[0]
            active_super_admins=check.execute("SELECT COUNT(*) FROM admin_users WHERE role='super_admin' AND is_active=1").fetchone()[0]
            if admin_count>0 and active_super_admins==0:
                raise ValueError("restored_database_has_no_active_super_admin")
            integrity = check.execute("PRAGMA integrity_check").fetchone()[0]
            foreign_key_errors = check.execute("PRAGMA foreign_key_check").fetchall()
            if integrity != "ok":
                raise ValueError("restored_database_integrity_check_failed")
            if foreign_key_errors:
                raise ValueError("restored_database_foreign_key_check_failed")

        live = Path(db.path)
        live.parent.mkdir(parents=True, exist_ok=True)

        # Coordinate across all M-Group workers. Database.connect() holds a
        # shared OS lock until each request closes its connection; the exclusive
        # lock below waits for in-flight work and prevents new DB connections
        # while the live file is switched.
        with db.exclusive_lock():
            if live.exists():
                # Flush any prior WAL content and switch the old database back
                # to rollback-journal mode before replacing its main file. This
                # prevents a stale -wal/-shm sidecar from being replayed against
                # the restored database after the rename.
                current = sqlite3.connect(live, timeout=30)
                try:
                    checkpoint = current.execute("PRAGMA wal_checkpoint(TRUNCATE)").fetchone()
                    if checkpoint and checkpoint[0] != 0:
                        raise ValueError("active_database_checkpoint_busy")
                    journal_mode = current.execute("PRAGMA journal_mode=DELETE").fetchone()
                    if not journal_mode or str(journal_mode[0]).lower() != "delete":
                        raise ValueError("active_database_journal_mode_switch_failed")
                finally:
                    current.close()

            # Sidecars are safe to remove only after all managed connections
            # have drained and the old DB has left WAL mode.
            for suffix in ("-wal", "-shm", "-journal"):
                Path(str(live) + suffix).unlink(missing_ok=True)
            tmp.chmod(0o600)
            tmp.replace(live)
    except Exception:
        if probe is not None:
            probe.close()
        tmp.unlink(missing_ok=True)
        Path(str(tmp) + "-journal").unlink(missing_ok=True)
        Path(str(tmp) + "-wal").unlink(missing_ok=True)
        Path(str(tmp) + "-shm").unlink(missing_ok=True)
        Path(str(tmp) + ".lock").unlink(missing_ok=True)
        raise
    Path(str(tmp) + ".lock").unlink(missing_ok=True)
    return manifest

def import_excel(db, source):
    if load_workbook is None:
        raise RuntimeError("openpyxl_required")
    wb=load_workbook(source,data_only=True,read_only=True)
    imported={}
    with db.connect() as c:
        if "cashier" in wb.sheetnames:
            ws=wb["cashier"]
            headers=[x.value for x in next(ws.iter_rows())]
            idx={h:i for i,h in enumerate(headers)}
            for row in ws.iter_rows(min_row=2,values_only=True):
                tid=row[idx["transaction_id"]]
                if not tid: continue
                c.execute("""INSERT OR IGNORE INTO cashier_deposits
                    (transaction_id,hardware_id,credential_id,member_id,account_id,amount,status,
                     balance_before,balance_after,created_at,completed_at)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                    (tid,row[idx["hardware_id"]],row[idx["credential_id"]],row[idx["member_id"]],
                     row[idx["account_id"]],int(row[idx["amount"]]),row[idx["status"]],
                     row[idx["balance_before"]],row[idx["balance_after"]],row[idx["created_at"]],row[idx["created_at"]]))
            imported["cashier"]=c.total_changes
        c.commit()
    return imported


def export_rows_excel(rows, target, metadata=None):
    if Workbook is None:
        raise RuntimeError("openpyxl_required")
    from datetime import datetime, timezone
    from zoneinfo import ZoneInfo
    from openpyxl.styles import Font, PatternFill, Alignment
    target=Path(target)
    target.parent.mkdir(parents=True,exist_ok=True)
    metadata=metadata or {}
    wb=Workbook()
    ws=wb.active
    ws.title="Laporan Transaksi"
    source_label={"ALL":"Semua perangkat","CASHIER":"Cashier","VENDING":"Vending","TOPUP":"Top Up"}.get(metadata.get("source"),metadata.get("source","Semua perangkat"))
    ws["A1"]="M-GROUP — LAPORAN TRANSAKSI"
    ws["A1"].font=Font(bold=True,size=16)
    ws["A2"]="Sumber: "+str(source_label)
    ws["A3"]="Periode: "+str(metadata.get("start","—"))+" sampai "+str(metadata.get("end","—"))+" (Asia/Jakarta)"
    ws["A4"]="Metode pembayaran: "+str(metadata.get("method","ALL"))
    ws["A5"]="Perangkat: "+str(metadata.get("hardware_id","ALL"))
    ws["A6"]="Dibuat: "+datetime.now(ZoneInfo("Asia/Jakarta")).strftime("%Y-%m-%d %H:%M:%S")+" WIB"
    rows=list(rows or [])
    completed=[r for r in rows if str(r.get("status","")).lower()=="completed"]
    failed=[r for r in rows if str(r.get("status","")).lower() in {"failed","error"}]
    canceled=[r for r in rows if str(r.get("status","")).lower() in {"canceled","cancelled","refunded"}]
    ws["A7"]="Ringkasan"
    ws["A6"].font=Font(bold=True)
    ws.append(["Status","Jumlah transaksi","Total nominal"])
    for cell in ws[8]:
        cell.font=Font(bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="1F4E78")
    for label,group in [("Berhasil",completed),("Gagal",failed),("Dibatalkan/Refund",canceled)]:
        ws.append([label,len(group),sum(int(r.get("amount") or 0) for r in group)])
    ws["A13"]="Detail transaksi"
    ws["A12"].font=Font(bold=True,size=12)
    headers=list(rows[0].keys()) if rows else ["transaction_id","created_at","source","hardware_id","member_id","account_id","transaction_type","amount","status","payment_method","payment_provider"]
    ws.append(headers)
    header_row=14
    for cell in ws[header_row]:
        cell.font=Font(bold=True,color="FFFFFF")
        cell.fill=PatternFill("solid",fgColor="1F4E78")
        cell.alignment=Alignment(wrap_text=True)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    ws.freeze_panes="A15"
    ws.auto_filter.ref=f"A{header_row}:{ws.cell(ws.max_row, len(headers)).coordinate}"
    for col in ws.columns:
        letter=col[0].column_letter
        width=min(34,max(12,max((len(str(cell.value or "")) for cell in col),default=10)+2))
        ws.column_dimensions[letter].width=width
    for row in ws.iter_rows(min_row=9,max_row=11,min_col=3,max_col=3):
        for cell in row: cell.number_format='"Rp" #,##0'
    for row in ws.iter_rows(min_row=header_row+1,min_col=1,max_col=len(headers)):
        for cell in row:
            if headers[cell.column-1]=="amount" and isinstance(cell.value,(int,float)):
                cell.number_format='"Rp" #,##0'
    wb.save(target)
    return target
