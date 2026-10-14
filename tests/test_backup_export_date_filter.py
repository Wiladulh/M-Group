from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from app.core import backup
from app.core.db import Database


def test_vending_excel_export_handles_joined_created_at_and_date_filter(monkeypatch):
    with TemporaryDirectory() as directory:
        root = Path(directory)
        db = Database(root / "m-group.db")
        with db.connect() as c:
            c.execute("""
                CREATE TABLE vending_machines (
                    machine_id TEXT PRIMARY KEY,
                    name TEXT NOT NULL,
                    created_at TEXT NOT NULL
                )
            """)
            c.execute("""
                CREATE TABLE vending_transactions (
                    transaction_id TEXT PRIMARY KEY,
                    machine_id TEXT NOT NULL,
                    amount INTEGER NOT NULL,
                    created_at TEXT NOT NULL
                )
            """)
            c.execute(
                "INSERT INTO vending_machines(machine_id,name,created_at) VALUES(?,?,?)",
                ("vm-1", "Machine", "2026-01-01 00:00:00"),
            )
            c.executemany(
                "INSERT INTO vending_transactions(transaction_id,machine_id,amount,created_at) VALUES(?,?,?,?)",
                [
                    ("before", "vm-1", 1000, "2025-12-31 23:59:59"),
                    ("inside", "vm-1", 2000, "2026-01-01 12:00:00"),
                    ("after", "vm-1", 3000, "2026-01-02 00:00:00"),
                ],
            )
            c.commit()

        monkeypatch.setattr(
            backup,
            "SHEETS",
            {
                "vending": """
                    SELECT v.transaction_id,v.created_at,m.machine_id,v.amount
                    FROM vending_transactions v
                    LEFT JOIN vending_machines m ON m.machine_id=v.machine_id
                    ORDER BY v.created_at
                """
            },
        )

        all_rows_path = root / "all.xlsx"
        backup.export_excel(db, all_rows_path)
        workbook = load_workbook(all_rows_path, read_only=True, data_only=True)
        try:
            all_rows = list(workbook["vending"].iter_rows(values_only=True))
        finally:
            workbook.close()
        assert [row[0] for row in all_rows[1:]] == ["before", "inside", "after"]

        filtered_path = root / "filtered.xlsx"
        backup.export_excel(db, filtered_path, "2026-01-01", "2026-01-02")
        workbook = load_workbook(filtered_path, read_only=True, data_only=True)
        try:
            filtered_rows = list(workbook["vending"].iter_rows(values_only=True))
        finally:
            workbook.close()
        assert [row[0] for row in filtered_rows[1:]] == ["inside"]
